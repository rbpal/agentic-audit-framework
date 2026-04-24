"""Walker that replaces ``<placeholder>`` markers with seeded fake data.

Consumes the output of ``render_toc_sheet(spec)`` (Task 2) and mutates
it in place, replacing every ``<placeholder>`` cell with a seeded fake
value. Uses a LOCAL ``random.Random(spec.seed)`` — never touches the
global ``random`` module state.
"""

from __future__ import annotations

import random
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell

from agentic_audit.generator import fake_data
from agentic_audit.models.scenario import ScenarioSpec

# NOTE: placeholder charset must be \w (A-Za-z0-9_) — the lowercase-only
# character class [a-z_0-9] silently rejects uppercase-bearing markers like
# <attribute_A_description>, leaving them unresolved in the output. The
# dispatch regexes below use [A-F] and expect this outer charset to admit
# uppercase letters. See tests/unit/generator/test_populate.py coverage
# for the regression case.
_PLACEHOLDER_RE = re.compile(r"^<(\w+)>$")
_EMBEDDED_PLACEHOLDER_RE = re.compile(r"<(\w+)>")
_SAMPLE_PLACEHOLDER_RE = re.compile(r"^sample_(\d+)_(\w+)$")
_ATTRIBUTE_PLACEHOLDER_RE = re.compile(r"^attribute_([A-F])_(\w+)$")


def populate_workbook(wb: Workbook, spec: ScenarioSpec) -> Workbook:
    """Replace every ``<placeholder>`` with a seeded fake value.

    Mutates ``wb`` in place; returns for chaining convenience.
    Raises ``ValueError`` on any unrecognized placeholder name — the
    guarantee is that no ``<xxx>`` markers remain after a successful call.

    Whole-cell markers (``<foo>``) preserve the resolver's native return
    type (e.g. ``datetime.date``, ``int``). Markers embedded in a
    compound string (``"DC-9: <control_name>"``) are substituted via
    ``str()`` since they must coexist with surrounding literal text.
    """
    rng = random.Random(spec.seed)
    ws = wb.active
    assert ws is not None

    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or "<" not in cell.value:
                continue
            whole = _PLACEHOLDER_RE.match(cell.value)
            if whole is not None:
                cell.value = _resolve_placeholder(whole.group(1), cell, rng, spec)
                continue
            cell.value = _EMBEDDED_PLACEHOLDER_RE.sub(
                lambda m: str(_resolve_placeholder(m.group(1), cell, rng, spec)),  # noqa: B023 — re.sub is synchronous; cell is stable for the call
                cell.value,
            )

    return wb


def _resolve_placeholder(name: str, cell: Cell, rng: random.Random, spec: ScenarioSpec) -> Any:  # noqa: ANN401  # openpyxl cell values are heterogeneous by design
    """Dispatch a placeholder name to its fake-data provider.

    Three levels of dispatch:
      1. Exact-match dict of simple providers
      2. Regex match for indexed placeholders (sample_N_X, attribute_L_X)
      3. Raise on unknown
    """
    # ── Level 1: exact-match simple providers
    simple = _SIMPLE_DISPATCH.get(name)
    if simple is not None:
        return simple(rng, spec, cell)

    # ── Level 2: indexed placeholders
    sample_match = _SAMPLE_PLACEHOLDER_RE.match(name)
    if sample_match:
        index = int(sample_match.group(1))
        kind = sample_match.group(2)
        if kind == "description":
            return fake_data.fake_sample_description(rng, spec, index)
        if kind == "period":
            return fake_data.fake_sample_period(rng, spec, index)
        if kind == "wp_ref":
            return fake_data.fake_wp_ref(rng)
        raise ValueError(f"Unknown sample placeholder kind: {name!r}")

    attribute_match = _ATTRIBUTE_PLACEHOLDER_RE.match(name)
    if attribute_match:
        letter = attribute_match.group(1)
        kind = attribute_match.group(2)
        if kind == "description":
            return fake_data.fake_attribute_description(spec, letter)
        if kind == "toc_procedure":
            return fake_data.fake_toc_procedure(spec, letter)
        raise ValueError(f"Unknown attribute placeholder kind: {name!r}")

    # ── Level 3: unknown
    raise ValueError(f"Unrecognized placeholder: <{name}>")


# ── Simple-provider dispatch table ───────────────────────────────────
# Every entry: placeholder name → callable (rng, spec, cell) → value.


def _workpaper_ref(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_workpaper_ref(rng)


def _entity_name(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_entity_name(rng)


def _scot_name(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_scot_name(rng)


def _gaas(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_gaas(rng)


def _engagement_nature(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_engagement_nature(rng)


def _year_end_date(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> Any:  # noqa: ANN401
    del cell
    return fake_data.fake_year_end_date(rng, spec.quarter)


def _effectiveness_conclusion(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del rng, cell
    return fake_data.effectiveness_conclusion(spec.expected_outcome)


def _preparer_initials(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_initials(rng)


def _control_description(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del cell
    return fake_data.fake_control_description(rng, spec)


def _control_type(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_control_type(rng)


def _yes_no(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_yes_no(rng)


def _it_applications(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_it_applications(rng)


def _population_description(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del cell
    return fake_data.fake_population_description(rng, spec)


def _exceptions_noted(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del rng, cell
    return fake_data.exceptions_noted(spec.expected_outcome)


def _exception_narrative(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del rng, cell
    return fake_data.fake_exception_narrative(spec)


def _compensating_names(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del rng, spec, cell
    return "N/A"


def _deficiency_refs(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del rng, spec, cell
    return "N/A"


def _sample_extension(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del rng, spec, cell
    return "N/A"


def _tickmark(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    # Column H (8) → A, I (9) → B, ... M (13) → F
    attribute_letter = chr(ord("A") + (cell.column - 8))
    return fake_data.fake_tickmark(rng, spec, attribute_letter)


def _scot_form_ref(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_workpaper_ref(rng)


def _ipe_risk_ref(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del spec, cell
    return fake_data.fake_workpaper_ref(rng)


def _control_name(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    """The title-row value is ``"{control_id}: <control_name>"``; only the
    ``<control_name>`` marker shows up alone when it has its own cell.
    """
    del rng, cell
    return (
        "Real-estate billing calculation"
        if spec.pattern_type == "signoff_with_tieout"
        else "Monthly revenue variance analysis"
    )


def _report_validation(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> str:
    del rng, spec, cell
    return "Configurations validated per IT general-controls workpaper."


def _toc_billing_fee_claim(rng: random.Random, spec: ScenarioSpec, cell: Cell) -> int:
    """DC-9 only (Task 13) — the billing fee the TOC *asserts*.

    For pass scenarios → equals the billing-calc W/P's canonical fee
    (the TOC correctly records what the supporting schedule shows).

    For ``figure_mismatch`` → intentionally offset from the canonical
    fee, producing a cross-file contradiction the agent must detect.
    """
    del rng, cell
    return fake_data.compute_toc_billing_claim(spec)


_SIMPLE_DISPATCH: dict[str, Any] = {
    "workpaper_ref": _workpaper_ref,
    "scot_form_ref": _scot_form_ref,
    "ipe_risk_ref": _ipe_risk_ref,
    "entity_name": _entity_name,
    "scot_name": _scot_name,
    "gaas": _gaas,
    "engagement_nature": _engagement_nature,
    "year_end_date": _year_end_date,
    "effectiveness_conclusion": _effectiveness_conclusion,
    "preparer_initials": _preparer_initials,
    "reviewer_1_initials": _preparer_initials,
    "reviewer_2_initials": _preparer_initials,
    "control_description": _control_description,
    "control_type": _control_type,
    "review_or_monitoring": _yes_no,
    "it_applications": _it_applications,
    "ipe_yes_no": _yes_no,
    "configurable_yes_no": _yes_no,
    "compensating_yes_no": _yes_no,
    "report_validation": _report_validation,
    "population_description": _population_description,
    "exceptions_noted": _exceptions_noted,
    "exceptions_random": _yes_no,
    "exception_nature": _exception_narrative,
    "sample_extension": _sample_extension,
    "compensating_names": _compensating_names,
    "deficiency_refs": _deficiency_refs,
    "tickmark": _tickmark,
    "control_name": _control_name,
    "toc_billing_fee_claim": _toc_billing_fee_claim,
}
